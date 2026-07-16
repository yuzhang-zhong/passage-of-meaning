from pathlib import Path

from openpyxl import load_workbook


path = Path("research/dict_revised_2015_20260625.xlsx")
workbook = load_workbook(path, read_only=True, data_only=True)

print({"sheets": workbook.sheetnames})
for sheet in workbook.worksheets:
    print({"sheet": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column})
    for row in sheet.iter_rows(min_row=1, max_row=min(6, sheet.max_row), values_only=True):
        print(row)
