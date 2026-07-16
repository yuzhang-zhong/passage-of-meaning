from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DICT_PATH = ROOT / "research" / "dict_revised_2015_20260625.xlsx"
CATALOG_PATH = ROOT / "src" / "catalog.ts"
OUTPUT_PATH = ROOT / "src" / "dictionaryEvidence.generated.json"
SOURCE_URL = "https://language.moe.gov.tw/001/Upload/Files/site_content/M0001/respub/dict_reviseddict_download.html"


def load_conversion_table(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    with path.open("r", encoding="utf-8") as source:
        for raw_line in source:
            line = raw_line.strip()
            if not line or "\t" not in line:
                continue
            simplified, traditional = line.split("\t", 1)
            result[simplified] = traditional.split(" ", 1)[0]
    return result


character_map = load_conversion_table(ROOT / "research" / "STCharacters.txt")
phrase_map = load_conversion_table(ROOT / "research" / "STPhrases.txt")


def to_traditional(text: str) -> str:
    if text in phrase_map:
        return phrase_map[text]
    return "".join(character_map.get(character, character) for character in text)


catalog_source = CATALOG_PATH.read_text(encoding="utf-8")
catalog_words: list[str] = []
for block in re.findall(r"words:\s*\[(.*?)\]", catalog_source, flags=re.S):
    catalog_words.extend(re.findall(r"'([^']+)'", block))
catalog_words = list(dict.fromkeys(catalog_words))

workbook = load_workbook(DICT_PATH, read_only=True, data_only=True)
sheet = workbook[workbook.sheetnames[0]]
entries: dict[str, list[str]] = defaultdict(list)

for headword, alias, *_middle, definition, _reference, _variant in sheet.iter_rows(min_row=2, values_only=True):
    if not isinstance(headword, str) or not isinstance(definition, str):
        continue
    entries[headword.strip()].append(definition)
    if isinstance(alias, str):
        for item in re.split(r"[、,，;；\s]+", alias.strip()):
            if item:
                entries[item].append(definition)


def clean_definition(raw: str) -> str:
    text = raw.replace("_x000D_", "\n")
    text = re.sub(r"\[[^\]]+\]", "", text)
    text = re.sub(r"^\s*\d+[.．]\s*", "", text, flags=re.M)
    text = re.sub(r"\s+", " ", text).strip()
    first = re.split(r"(?<=[。！？])", text, maxsplit=1)[0].strip()
    if len(first) < 12:
        first = text[:180].strip()
    return first[:180].rstrip()


evidence: dict[str, dict[str, object]] = {}
missing: list[dict[str, str]] = []

for word in catalog_words:
    traditional = to_traditional(word)
    definitions = entries.get(traditional) or entries.get(word)
    if not definitions:
        missing.append({"word": word, "query": traditional})
        continue
    definition = max(definitions, key=len)
    excerpt = clean_definition(definition)
    evidence[word] = {
        "id": f"moe-dict-{word}",
        "excerpt": f"“{word}”條：{excerpt}",
        "context": f"教育部《重編國語辭典修訂本》將“{traditional}”收為辭目。這張卡只用來核對辭書所記的詞義邊界，不據此推定該詞義的首見年份。",
        "source": f"教育部《重編國語辭典修訂本》·“{traditional}”條",
        "year": 2026,
        "dateLabel": "網路第六版 · 2026-06-25",
        "medium": "歷史語言辭典",
        "grade": "A",
        "status": "人文解釋",
        "note": "官方授權資料庫中的辭目與釋義；辭書收錄可證用法存在，不等於首次出現年代。",
        "sourceUrl": SOURCE_URL,
        "reviewed": "已核驗",
    }

OUTPUT_PATH.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({
    "catalogWords": len(catalog_words),
    "matched": len(evidence),
    "missing": len(missing),
    "missingWords": missing,
    "output": str(OUTPUT_PATH),
}, ensure_ascii=False, indent=2))
